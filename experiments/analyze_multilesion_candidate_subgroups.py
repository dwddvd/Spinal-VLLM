#!/usr/bin/env python
"""Patient-level multi-candidate and clinical multi-lesion subgroup analysis.

This is a post-processing analysis. It does not run nnU-Net or Qwen and does
not require a GPU.

The inference-derived subgroup is deliberately called a "multi-candidate
pattern", not clinical multifocal disease. A clinical single/multiple-lesion
subgroup is analyzed only when an independently reviewed patient-level CSV is
provided with --clinical_labels_csv.
"""

import argparse
import csv
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path


CLASS_LABELS = ("infection", "tumor")
DEFAULT_STRATEGIES = {
    "weighted_top3": "current_weighted_top3_fallback_r3_patient_predictions.csv",
    "top1": "top1_component_fallback_r3_patient_predictions.csv",
    "component_majority": "component_majority_fallback_r3_patient_predictions.csv",
}
SUBGROUP_DEFINITIONS = (
    "candidate_any_multi",
    "candidate_ge2_slices",
    "candidate_frac20",
)
METRIC_NAMES = (
    "accuracy",
    "balanced_accuracy",
    "infection_accuracy",
    "tumor_accuracy",
)


def safe_int(value, default=0):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def safe_div(numerator, denominator):
    return numerator / denominator if denominator else 0.0


def percentile(values, q):
    if not values:
        return 0.0
    values = sorted(values)
    position = (len(values) - 1) * q
    lower = int(math.floor(position))
    upper = int(math.ceil(position))
    if lower == upper:
        return values[lower]
    fraction = position - lower
    return values[lower] * (1.0 - fraction) + values[upper] * fraction


def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows, fieldnames=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def patient_id_from_row(row):
    patient_id = str(row.get("patient_id", "")).strip()
    if patient_id:
        return patient_id
    parts = str(row.get("case_id", "")).split("_")
    return parts[1] if len(parts) > 1 else ""


def gt_label_from_rows(rows):
    labels = {
        str(row.get("gt_label", "")).strip().lower()
        for row in rows
        if str(row.get("gt_label", "")).strip()
    }
    if len(labels) != 1:
        raise ValueError(f"Inconsistent or missing patient GT labels: {sorted(labels)}")
    label = next(iter(labels))
    if label not in CLASS_LABELS:
        raise ValueError(f"Unsupported GT label: {label}")
    return label


def build_candidate_burden(slice_rows):
    grouped = defaultdict(list)
    for row in slice_rows:
        patient_id = patient_id_from_row(row)
        if not patient_id:
            raise ValueError(f"Unable to infer patient_id from row: {row}")
        grouped[patient_id].append(row)

    records = []
    for patient_id, rows in sorted(grouped.items()):
        valid_votes = [safe_int(row.get("valid_component_votes")) for row in rows]
        raw_components = [safe_int(row.get("num_raw_components")) for row in rows]
        candidate_slices = sum(value >= 1 for value in valid_votes)
        multi_slices = sum(value >= 2 for value in valid_votes)
        sequences = sorted({
            str(row.get("seq", "")).strip().upper()
            for row in rows
            if str(row.get("seq", "")).strip()
        })
        records.append({
            "patient_id": patient_id,
            "gt_label": gt_label_from_rows(rows),
            "sequences": ",".join(sequences),
            "num_sequences": len(sequences),
            "num_slice_rows": len(rows),
            "num_candidate_slices": candidate_slices,
            "num_no_candidate_slices": len(rows) - candidate_slices,
            "num_multicandidate_slices": multi_slices,
            "multicandidate_slice_fraction": safe_div(multi_slices, candidate_slices),
            "max_valid_component_votes": max(valid_votes, default=0),
            "max_raw_components": max(raw_components, default=0),
            "num_slices_raw_components_ge2": sum(value >= 2 for value in raw_components),
            "candidate_any_multi": int(multi_slices >= 1),
            "candidate_ge2_slices": int(multi_slices >= 2),
            "candidate_frac20": int(
                candidate_slices > 0
                and safe_div(multi_slices, candidate_slices) >= 0.20
            ),
        })
    return records


def normalize_clinical_pattern(value):
    value = str(value or "").strip().lower().replace("-", "_").replace(" ", "_")
    single_values = {"single", "single_lesion", "unifocal", "solitary", "0"}
    multiple_values = {
        "multiple",
        "multiple_lesions",
        "multifocal",
        "multi_lesion",
        "1",
    }
    if value in single_values:
        return "single"
    if value in multiple_values:
        return "multiple"
    return ""


def load_clinical_labels(path):
    rows = read_csv(path)
    labels = {}
    for row in rows:
        patient_id = str(row.get("patient_id", "")).strip()
        if not patient_id:
            continue
        pattern = normalize_clinical_pattern(
            row.get("lesion_pattern") or row.get("clinical_lesion_pattern")
        )
        if patient_id in labels and labels[patient_id] != pattern:
            raise ValueError(f"Duplicate conflicting clinical label: {patient_id}")
        labels[patient_id] = pattern
    return labels


def load_strategy_predictions(strategy_dir):
    strategy_dir = Path(strategy_dir)
    outputs = {}
    for strategy, filename in DEFAULT_STRATEGIES.items():
        path = strategy_dir / filename
        if not path.exists():
            raise FileNotFoundError(f"Missing strategy prediction file: {path}")
        rows = read_csv(path)
        by_patient = {}
        for row in rows:
            patient_id = str(row.get("patient_id", "")).strip()
            if not patient_id:
                raise ValueError(f"Missing patient_id in {path}")
            if patient_id in by_patient:
                raise ValueError(f"Duplicate patient_id {patient_id} in {path}")
            gt_label = str(row.get("gt_label", "")).strip().lower()
            pred_label = str(row.get("pred_label", "")).strip().lower()
            by_patient[patient_id] = {
                **row,
                "patient_id": patient_id,
                "gt_label": gt_label,
                "pred_label": pred_label,
                "is_correct": int(gt_label == pred_label),
            }
        outputs[strategy] = by_patient
    common = set.intersection(*(set(rows) for rows in outputs.values()))
    if not common:
        raise ValueError("No common patients across strategy prediction files.")
    for strategy, rows in outputs.items():
        if set(rows) != common:
            raise ValueError(
                f"Patient set differs for {strategy}: "
                f"{len(rows)} patients versus {len(common)} common"
            )
    return outputs


def metrics_for_rows(rows):
    total = len(rows)
    correct = sum(row["gt_label"] == row["pred_label"] for row in rows)
    by_label = {}
    for label in CLASS_LABELS:
        label_rows = [row for row in rows if row["gt_label"] == label]
        label_correct = sum(row["pred_label"] == label for row in label_rows)
        by_label[label] = {
            "n": len(label_rows),
            "correct": label_correct,
            "accuracy": safe_div(label_correct, len(label_rows)),
        }
    return {
        "n": total,
        "correct": correct,
        "accuracy": safe_div(correct, total),
        "balanced_accuracy": (
            by_label["infection"]["accuracy"] + by_label["tumor"]["accuracy"]
        ) / 2.0,
        "infection_n": by_label["infection"]["n"],
        "infection_accuracy": by_label["infection"]["accuracy"],
        "tumor_n": by_label["tumor"]["n"],
        "tumor_accuracy": by_label["tumor"]["accuracy"],
        "confusion": dict(Counter(
            f'{row["gt_label"]}->{row["pred_label"]}' for row in rows
        )),
    }


def stratified_bootstrap_ids(patient_ids, patient_gt, rng):
    sampled = []
    for label in CLASS_LABELS:
        label_ids = [patient_id for patient_id in patient_ids if patient_gt[patient_id] == label]
        sampled.extend(rng.choice(label_ids) for _ in label_ids)
    return sampled


def bootstrap_subgroup(
    patient_ids,
    strategy_predictions,
    n_bootstrap,
    seed,
):
    patient_gt = {
        patient_id: strategy_predictions["weighted_top3"][patient_id]["gt_label"]
        for patient_id in patient_ids
    }
    rng = random.Random(seed)
    samples = []
    for _ in range(n_bootstrap):
        sampled_ids = stratified_bootstrap_ids(patient_ids, patient_gt, rng)
        sample = {}
        for strategy, predictions in strategy_predictions.items():
            rows = [predictions[patient_id] for patient_id in sampled_ids]
            point = metrics_for_rows(rows)
            for metric in METRIC_NAMES:
                sample[f"{strategy}__{metric}"] = point[metric]
        for comparator in ("top1", "component_majority"):
            for metric in METRIC_NAMES:
                sample[f"delta_weighted_top3_minus_{comparator}__{metric}"] = (
                    sample[f"weighted_top3__{metric}"]
                    - sample[f"{comparator}__{metric}"]
                )
        samples.append(sample)

    summary = {}
    for key in samples[0]:
        values = [sample[key] for sample in samples]
        summary[key] = {
            "bootstrap_mean": sum(values) / len(values),
            "ci95_low": percentile(values, 0.025),
            "ci95_high": percentile(values, 0.975),
        }
    return summary


def bootstrap_interaction(
    positive_ids,
    negative_ids,
    strategy_predictions,
    n_bootstrap,
    seed,
):
    patient_gt = {
        patient_id: strategy_predictions["weighted_top3"][patient_id]["gt_label"]
        for patient_id in positive_ids + negative_ids
    }
    rng = random.Random(seed)
    distributions = defaultdict(list)
    for _ in range(n_bootstrap):
        sampled_positive = stratified_bootstrap_ids(positive_ids, patient_gt, rng)
        sampled_negative = stratified_bootstrap_ids(negative_ids, patient_gt, rng)
        group_metrics = {}
        for group_name, sampled_ids in (
            ("positive", sampled_positive),
            ("negative", sampled_negative),
        ):
            group_metrics[group_name] = {}
            for strategy, predictions in strategy_predictions.items():
                rows = [predictions[patient_id] for patient_id in sampled_ids]
                group_metrics[group_name][strategy] = metrics_for_rows(rows)

        for comparator in ("top1", "component_majority"):
            for metric in METRIC_NAMES:
                positive_delta = (
                    group_metrics["positive"]["weighted_top3"][metric]
                    - group_metrics["positive"][comparator][metric]
                )
                negative_delta = (
                    group_metrics["negative"]["weighted_top3"][metric]
                    - group_metrics["negative"][comparator][metric]
                )
                key = f"weighted_top3_minus_{comparator}__{metric}"
                distributions[key].append(positive_delta - negative_delta)

    summary = {}
    for key, values in distributions.items():
        summary[key] = {
            "bootstrap_mean": sum(values) / len(values),
            "ci95_low": percentile(values, 0.025),
            "ci95_high": percentile(values, 0.975),
        }
    return summary


def exact_mcnemar_p(reference_only_correct, comparator_only_correct):
    n = reference_only_correct + comparator_only_correct
    if n == 0:
        return 1.0
    lower = min(reference_only_correct, comparator_only_correct)
    tail = sum(math.comb(n, k) for k in range(lower + 1)) / (2 ** n)
    return min(1.0, 2.0 * tail)


def subgroup_members(patient_records, definition):
    included = []
    positive = []
    negative = []
    for record in patient_records:
        value = record.get(definition, "")
        if value in ("", None):
            continue
        included.append(record["patient_id"])
        if safe_int(value) == 1:
            positive.append(record["patient_id"])
        else:
            negative.append(record["patient_id"])
    return included, positive, negative


def analyze_definition(
    definition,
    positive_name,
    negative_name,
    patient_records,
    strategy_predictions,
    n_bootstrap,
    seed,
):
    included, positive_ids, negative_ids = subgroup_members(patient_records, definition)
    if not positive_ids or not negative_ids:
        return [], [], [], {
            "definition": definition,
            "status": "skipped",
            "reason": "One subgroup is empty.",
            "num_included": len(included),
            "num_positive": len(positive_ids),
            "num_negative": len(negative_ids),
        }

    metric_rows = []
    bootstrap_rows = []
    subgroup_map = {
        positive_name: positive_ids,
        negative_name: negative_ids,
    }
    bootstrap_by_group = {}

    for subgroup, patient_ids in subgroup_map.items():
        bootstrap_by_group[subgroup] = bootstrap_subgroup(
            patient_ids,
            strategy_predictions,
            n_bootstrap=n_bootstrap,
            seed=seed + (0 if subgroup == positive_name else 100003),
        )
        for strategy, predictions in strategy_predictions.items():
            point = metrics_for_rows([predictions[patient_id] for patient_id in patient_ids])
            row = {
                "definition": definition,
                "subgroup": subgroup,
                "strategy": strategy,
                **{key: value for key, value in point.items() if key != "confusion"},
                "confusion": json.dumps(point["confusion"], ensure_ascii=False, sort_keys=True),
            }
            for metric in METRIC_NAMES:
                ci = bootstrap_by_group[subgroup][f"{strategy}__{metric}"]
                row[f"{metric}_ci95_low"] = ci["ci95_low"]
                row[f"{metric}_ci95_high"] = ci["ci95_high"]
            metric_rows.append(row)

        for comparator in ("top1", "component_majority"):
            for metric in METRIC_NAMES:
                key = f"delta_weighted_top3_minus_{comparator}__{metric}"
                ci = bootstrap_by_group[subgroup][key]
                weighted_point = metrics_for_rows([
                    strategy_predictions["weighted_top3"][patient_id]
                    for patient_id in patient_ids
                ])[metric]
                comparator_point = metrics_for_rows([
                    strategy_predictions[comparator][patient_id]
                    for patient_id in patient_ids
                ])[metric]
                bootstrap_rows.append({
                    "definition": definition,
                    "subgroup": subgroup,
                    "comparison": f"weighted_top3_minus_{comparator}",
                    "metric": metric,
                    "point_delta": weighted_point - comparator_point,
                    "bootstrap_mean_delta": ci["bootstrap_mean"],
                    "ci95_low": ci["ci95_low"],
                    "ci95_high": ci["ci95_high"],
                })

    interaction_summary = bootstrap_interaction(
        positive_ids,
        negative_ids,
        strategy_predictions,
        n_bootstrap=n_bootstrap,
        seed=seed + 700001,
    )
    for comparator in ("top1", "component_majority"):
        for metric in METRIC_NAMES:
            key = f"weighted_top3_minus_{comparator}__{metric}"
            positive_weighted = metrics_for_rows([
                strategy_predictions["weighted_top3"][patient_id]
                for patient_id in positive_ids
            ])[metric]
            positive_comparator = metrics_for_rows([
                strategy_predictions[comparator][patient_id]
                for patient_id in positive_ids
            ])[metric]
            negative_weighted = metrics_for_rows([
                strategy_predictions["weighted_top3"][patient_id]
                for patient_id in negative_ids
            ])[metric]
            negative_comparator = metrics_for_rows([
                strategy_predictions[comparator][patient_id]
                for patient_id in negative_ids
            ])[metric]
            interaction = interaction_summary[key]
            bootstrap_rows.append({
                "definition": definition,
                "subgroup": f"interaction_{positive_name}_minus_{negative_name}",
                "comparison": f"weighted_top3_minus_{comparator}",
                "metric": metric,
                "point_delta": (
                    (positive_weighted - positive_comparator)
                    - (negative_weighted - negative_comparator)
                ),
                "bootstrap_mean_delta": interaction["bootstrap_mean"],
                "ci95_low": interaction["ci95_low"],
                "ci95_high": interaction["ci95_high"],
            })

    discordance_rows = []
    strategy_pairs = (
        ("weighted_top3", "top1"),
        ("weighted_top3", "component_majority"),
        ("top1", "component_majority"),
    )
    for subgroup, patient_ids in subgroup_map.items():
        for reference, comparator in strategy_pairs:
            both_correct = 0
            both_wrong = 0
            reference_only = 0
            comparator_only = 0
            different_predictions = 0
            for patient_id in patient_ids:
                ref = strategy_predictions[reference][patient_id]
                comp = strategy_predictions[comparator][patient_id]
                ref_correct = ref["is_correct"]
                comp_correct = comp["is_correct"]
                both_correct += int(ref_correct and comp_correct)
                both_wrong += int(not ref_correct and not comp_correct)
                reference_only += int(ref_correct and not comp_correct)
                comparator_only += int(not ref_correct and comp_correct)
                different_predictions += int(ref["pred_label"] != comp["pred_label"])
            discordance_rows.append({
                "definition": definition,
                "subgroup": subgroup,
                "reference": reference,
                "comparator": comparator,
                "n": len(patient_ids),
                "same_prediction": len(patient_ids) - different_predictions,
                "different_prediction": different_predictions,
                "both_correct": both_correct,
                "both_wrong": both_wrong,
                "reference_only_correct": reference_only,
                "comparator_only_correct": comparator_only,
                "exact_mcnemar_p": exact_mcnemar_p(reference_only, comparator_only),
            })

    status = {
        "definition": definition,
        "status": "completed",
        "num_included": len(included),
        "num_positive": len(positive_ids),
        "num_negative": len(negative_ids),
        "positive_name": positive_name,
        "negative_name": negative_name,
    }
    return metric_rows, bootstrap_rows, discordance_rows, status


def create_clinical_template(path, patient_records):
    rows = []
    for record in patient_records:
        rows.append({
            "patient_id": record["patient_id"],
            "gt_label": record["gt_label"],
            "lesion_pattern": "",
            "lesion_count": "",
            "multilevel_involvement": "",
            "reviewer_1": "",
            "reviewer_2": "",
            "adjudicated": "",
            "notes": "",
        })
    write_csv(path, rows)


def write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for XLSX output. Install it or rerun with --no_xlsx."
        ) from exc

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E8F3EA")

    for sheet_name, rows in sheets:
        ws = wb.create_sheet(title=sheet_name[:31])
        if not rows:
            ws.append(["No rows"])
            continue
        headers = list(rows[0])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_index, header in enumerate(headers, start=1):
            max_length = len(str(header))
            for cell in ws.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=min(ws.max_row, 300),
            ):
                for value_cell in cell:
                    max_length = max(max_length, len(str(value_cell.value or "")))
            ws.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 10),
                38,
            )
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            if "definition" in str(cell.value).lower():
                cell.fill = section_fill
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Analyze patient-level performance in inference-derived multi-candidate "
            "subgroups and, optionally, independently reviewed clinical multi-lesion "
            "subgroups."
        )
    )
    parser.add_argument("--slice_pred_csv", required=True)
    parser.add_argument("--strategy_dir", required=True)
    parser.add_argument("--output_dir", required=True)
    parser.add_argument("--clinical_labels_csv", default="")
    parser.add_argument("--model_name", default="")
    parser.add_argument("--n_bootstrap", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--no_xlsx", action="store_true")
    args = parser.parse_args()

    slice_rows = read_csv(Path(args.slice_pred_csv))
    if not slice_rows:
        raise ValueError(f"No slice prediction rows loaded from {args.slice_pred_csv}")
    patient_records = build_candidate_burden(slice_rows)
    strategy_predictions = load_strategy_predictions(args.strategy_dir)

    burden_ids = {row["patient_id"] for row in patient_records}
    strategy_ids = set(strategy_predictions["weighted_top3"])
    if burden_ids != strategy_ids:
        missing_burden = sorted(strategy_ids - burden_ids)
        missing_strategy = sorted(burden_ids - strategy_ids)
        raise ValueError(
            "Patient sets differ between slice predictions and strategy outputs. "
            f"Missing burden={missing_burden[:10]}, "
            f"missing strategy={missing_strategy[:10]}"
        )

    clinical_loaded = False
    clinical_missing = []
    if args.clinical_labels_csv:
        clinical_labels = load_clinical_labels(Path(args.clinical_labels_csv))
        clinical_loaded = True
        for record in patient_records:
            pattern = clinical_labels.get(record["patient_id"], "")
            record["clinical_lesion_pattern"] = pattern
            record["clinical_multiple"] = (
                1 if pattern == "multiple" else 0 if pattern == "single" else ""
            )
            if not pattern:
                clinical_missing.append(record["patient_id"])

    definitions = [
        ("candidate_any_multi", "multi_candidate", "single_candidate_only"),
        ("candidate_ge2_slices", "multi_candidate", "single_candidate_only"),
        ("candidate_frac20", "multi_candidate", "single_candidate_only"),
    ]
    if clinical_loaded:
        definitions.append(
            ("clinical_multiple", "clinical_multiple_lesions", "clinical_single_lesion")
        )

    metric_rows = []
    bootstrap_rows = []
    discordance_rows = []
    definition_status = []
    for index, (definition, positive_name, negative_name) in enumerate(definitions):
        metrics, bootstrap, discordance, status = analyze_definition(
            definition=definition,
            positive_name=positive_name,
            negative_name=negative_name,
            patient_records=patient_records,
            strategy_predictions=strategy_predictions,
            n_bootstrap=args.n_bootstrap,
            seed=args.seed + index * 1000003,
        )
        metric_rows.extend(metrics)
        bootstrap_rows.extend(bootstrap)
        discordance_rows.extend(discordance)
        definition_status.append(status)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    burden_csv = output_dir / "patient_candidate_burden.csv"
    metrics_csv = output_dir / "subgroup_metrics_long.csv"
    bootstrap_csv = output_dir / "subgroup_strategy_deltas_bootstrap_ci.csv"
    discordance_csv = output_dir / "subgroup_strategy_discordance.csv"
    template_csv = output_dir / "clinical_multilesion_review_template.csv"
    summary_json = output_dir / "multilesion_candidate_subgroup_summary.json"
    xlsx_path = output_dir / "multilesion_candidate_subgroup_analysis.xlsx"

    write_csv(burden_csv, patient_records)
    write_csv(metrics_csv, metric_rows)
    write_csv(bootstrap_csv, bootstrap_rows)
    write_csv(discordance_csv, discordance_rows)
    create_clinical_template(template_csv, patient_records)

    summary = {
        "analysis": "multilesion_and_multicandidate_patient_subgroups",
        "model_name": args.model_name,
        "slice_pred_csv": str(Path(args.slice_pred_csv)),
        "strategy_dir": str(Path(args.strategy_dir)),
        "clinical_labels_csv": str(Path(args.clinical_labels_csv)) if clinical_loaded else "",
        "num_slice_rows": len(slice_rows),
        "num_patients": len(patient_records),
        "patient_label_counts": dict(Counter(row["gt_label"] for row in patient_records)),
        "candidate_vote_counts": dict(Counter(
            str(safe_int(row.get("valid_component_votes"))) for row in slice_rows
        )),
        "definition_status": definition_status,
        "clinical_labels": {
            "loaded": clinical_loaded,
            "num_missing_or_unresolved": len(clinical_missing),
            "missing_patient_ids": clinical_missing,
        },
        "interpretation_guardrail": (
            "Inference-derived connected components/candidates are an engineering "
            "multi-candidate pattern and must not be described as clinical multifocal "
            "disease. Clinical multiple-lesion analysis is performed only from an "
            "independently reviewed patient-level lesion_pattern field."
        ),
        "pre_specified_definitions": {
            "primary": (
                "candidate_any_multi: at least one candidate-positive slice has two "
                "or more valid retained component candidates."
            ),
            "sensitivity_1": (
                "candidate_ge2_slices: at least two slices have two or more valid "
                "retained component candidates."
            ),
            "sensitivity_2": (
                "candidate_frac20: at least 20% of candidate-positive slices have "
                "two or more valid retained component candidates."
            ),
            "clinical": (
                "clinical_multiple: supplied by independent radiologist review; no "
                "clinical label is inferred from connected components."
            ),
        },
        "bootstrap": {
            "method": "patient-level stratified bootstrap within GT class",
            "n_bootstrap": args.n_bootstrap,
            "seed": args.seed,
        },
        "files": {
            "patient_candidate_burden_csv": str(burden_csv),
            "subgroup_metrics_csv": str(metrics_csv),
            "strategy_delta_bootstrap_csv": str(bootstrap_csv),
            "strategy_discordance_csv": str(discordance_csv),
            "clinical_review_template_csv": str(template_csv),
            "xlsx": "" if args.no_xlsx else str(xlsx_path),
        },
    }
    with open(summary_json, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    if not args.no_xlsx:
        definition_rows = []
        for item in definition_status:
            definition_rows.append(item)
        write_xlsx(
            xlsx_path,
            [
                ("Definitions", definition_rows),
                ("Patient burden", patient_records),
                ("Subgroup metrics", metric_rows),
                ("Strategy deltas", bootstrap_rows),
                ("Discordance", discordance_rows),
            ],
        )

    print(json.dumps({
        "summary_json": str(summary_json),
        "xlsx": "" if args.no_xlsx else str(xlsx_path),
        "num_patients": len(patient_records),
        "definition_status": definition_status,
        "clinical_labels_loaded": clinical_loaded,
        "clinical_labels_missing": len(clinical_missing),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
